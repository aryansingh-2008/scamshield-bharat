import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    # Institutional Palette
    WHITE = RGBColor(255, 255, 255)
    DARK = RGBColor(23, 32, 42)          # #17202A
    SLATE = RGBColor(71, 84, 103)        # #475467
    MUTED = RGBColor(102, 112, 133)      # #667085
    DEEP_BLUE = RGBColor(22, 78, 120)    # #164E78
    LIGHT_BLUE = RGBColor(240, 245, 250) # #F0F5FA
    BORDER_BLUE = RGBColor(197, 215, 229)
    
    RED_TEXT = RGBColor(180, 35, 24)     # #B42318
    RED_BG = RGBColor(254, 243, 242)     # #FEF3F2
    RED_BORDER = RGBColor(254, 205, 202) # #FECDCA

    AMBER_TEXT = RGBColor(181, 71, 8)    # #B54708
    AMBER_BG = RGBColor(255, 250, 235)   # #FFFAEB
    AMBER_BORDER = RGBColor(254, 223, 137)

    GREEN_TEXT = RGBColor(8, 122, 91)    # #087A5B
    GREEN_BG = RGBColor(240, 253, 244)   # #F0FDF4
    GREEN_BORDER = RGBColor(187, 247, 208)

    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BORDER = RGBColor(216, 222, 229)# #D8DEE5
    DIVIDER = RGBColor(226, 232, 240)    # #E2E8F0

    def set_white_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = WHITE

    def add_header(slide, tag, title, subtitle):
        set_white_background(slide)
        
        # Tag Badge
        tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.32))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = DEEP_BLUE
        tag_box.line.color.rgb = DEEP_BLUE
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

        # Title
        title_box = slide.shapes.add_textbox(Inches(3.15), Inches(0.35), Inches(9.38), Inches(0.45))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = DARK

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.73), Inches(0.35))
        tf = sub_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11.5)
        p.font.color.rgb = MUTED

        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.18), Inches(11.73), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = DIVIDER
        line.line.color.rgb = DIVIDER

    def add_footer(slide, slide_num):
        # Footer line
        fline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.73), Inches(0.012))
        fline.fill.solid()
        fline.fill.fore_color.rgb = DIVIDER
        fline.line.color.rgb = DIVIDER

        # Footer Text
        fbox = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(5.0), Inches(0.3))
        tf = fbox.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "ScamShield Bharat — Evidence-First Investor Safety Console"
        p.font.size = Pt(9)
        p.font.color.rgb = MUTED

        cbox = slide.shapes.add_textbox(Inches(5.0), Inches(6.95), Inches(4.5), Inches(0.3))
        tf = cbox.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "SANGYAN Hackathon (SNTC, IIT (BHU) × SEBI × NSDL)"
        p.font.size = Pt(9)
        p.font.color.rgb = MUTED
        p.alignment = PP_ALIGN.CENTER

        rbox = slide.shapes.add_textbox(Inches(10.0), Inches(6.95), Inches(2.53), Inches(0.3))
        tf = rbox.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Slide {slide_num} of 15"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = DEEP_BLUE
        p.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: Title & Executive Vision
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_white_background(slide1)

    # Top Brand Bar
    tbar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    tbar.fill.solid()
    tbar.fill.fore_color.rgb = DEEP_BLUE
    tbar.line.fill.background()

    # Hackathon Badge
    hbadge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(5.8), Inches(0.38))
    hbadge.fill.solid()
    hbadge.fill.fore_color.rgb = LIGHT_BLUE
    hbadge.line.color.rgb = BORDER_BLUE
    tf = hbadge.text_frame
    p = tf.paragraphs[0]
    p.text = "SANGYAN — INVESTOR RESILIENCE HACKATHON 2026"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p.alignment = PP_ALIGN.CENTER

    # Project Title
    tbox = slide1.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.73), Inches(1.1))
    tf = tbox.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "SCAMSHIELD BHARAT"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = DARK

    # Tagline Box
    tagbox = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.6), Inches(11.73), Inches(0.8))
    tagbox.fill.solid()
    tagbox.fill.fore_color.rgb = LIGHT_BLUE
    tagbox.line.color.rgb = BORDER_BLUE
    tagbox.line.width = Pt(1.5)
    tf = tagbox.text_frame
    tf.margin_left = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "“Check the claim before you act.”"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "Evidence-first investor safety console checking suspicious financial messages against trusted regulatory records."
    p2.font.size = Pt(12)
    p2.font.color.rgb = SLATE

    # 3 Metrics / Info Cards
    card_w = Inches(3.64)
    card_h = Inches(2.6)
    
    # Card 1: Collaboration
    c1 = create_card(slide1, Inches(0.8), Inches(3.65), card_w, card_h, WHITE, CARD_BORDER)
    tf = c1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "ORGANIZERS & PARTNERS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• SNTC, IIT (BHU) Varanasi\n• In collaboration with SEBI\n• Supported by NSDL\n• Theme: Investor Resilience"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Card 2: Live Prototype
    c2 = create_card(slide1, Inches(4.84), Inches(3.65), card_w, card_h, WHITE, CARD_BORDER)
    tf = c2.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "LIVE DEPLOYMENT & REPO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Production URL:\n  scamshield-bharat-kappa.vercel.app\n• Public GitHub Repository:\n  github.com/aryansingh-2008/\n  scamshield-bharat"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Card 3: Validation
    c3 = create_card(slide1, Inches(8.88), Inches(3.65), card_w, card_h, WHITE, CARD_BORDER)
    tf = c3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "ENGINEERING VALIDATION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GREEN_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• 79 / 79 Vitest Tests Passing\n• 54 / 54 Red-Team Vectors Cleared\n• Zero Disk Writes (In-Memory)\n• Bilingual Hindi & English Ready"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    add_footer(slide1, 1)

    # ==========================================
    # SLIDE 2: The Problem Space
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "PROBLEM STATEMENT", "The Moment Before Money Leaves", "Financial frauds exploit cognitive urgency and regulatory unawareness before victims can verify")
    
    # 4 Problem Cards (2x2 Grid)
    p_w = Inches(5.66)
    p_h = Inches(2.55)
    
    # Box 1
    b1 = create_card(slide2, Inches(0.8), Inches(1.4), p_w, p_h, RED_BG, RED_BORDER)
    tf = b1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "1. Guaranteed Return Tipping Scams"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• Fraudulent Telegram/WhatsApp VIP channels promise 300% monthly returns.\n• Fake screenshots of institutional Demat profits create FOMO.\n• Unregistered operators funnel victims into unregulated mule accounts."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Box 2
    b2 = create_card(slide2, Inches(6.86), Inches(1.4), p_w, p_h, RED_BG, RED_BORDER)
    tf = b2.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "2. Regulatory & Authority Impersonation"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• Forged SEBI certificates, RBI circulars, and police arrest warrants.\n• Threat of tax penalties, narcotics charges, and coercive 'Digital Arrest'.\n• Victims comply out of panic without accessible institutional verification."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Box 3
    b3 = create_card(slide2, Inches(0.8), Inches(4.15), p_w, p_h, AMBER_BG, AMBER_BORDER)
    tf = b3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "3. Fake KYC Expiry & Account Block Panics"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AMBER_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• Coercive SMS: 'Demat / Bank Account will be suspended tonight'.\n• Clone portal links (bit.ly / .xyz) capture netbanking passwords & OTPs.\n• Shortened URLs obscure phishing domain reputations."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Box 4
    b4 = create_card(slide2, Inches(6.86), Inches(4.15), p_w, p_h, RED_BG, RED_BORDER)
    tf = b4.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "4. Malicious APK & Screen Sharing Traps"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• Impersonators instruct users to sideload customized APK files.\n• Screen-sharing tools (AnyDesk, QuickSupport) capture banking OTPs.\n• Immediate unauthorized funds transfer from victims' accounts."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    add_footer(slide2, 2)

    # ==========================================
    # SLIDE 3: Who We Build For
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "TARGET PERSONAS", "Who We Build For: Protecting Every Citizen", "A first-time retail investor should never need to be a cybersecurity expert")

    card_w = Inches(3.64)
    card_h = Inches(4.3)

    p1 = create_card(slide3, Inches(0.8), Inches(1.4), card_w, card_h, WHITE, CARD_BORDER)
    tf = p1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "TIER-2 & TIER-3 INVESTORS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• First generation navigating smartphone Demat & trading apps.\n• Receives forwarded investment tips on WhatsApp & Telegram.\n• Needs simple, plain Hindi & English explanations with zero complex legal jargon.\n• High vulnerability to high-return lure scams."
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    p2_card = create_card(slide3, Inches(4.84), Inches(1.4), card_w, card_h, WHITE, CARD_BORDER)
    tf = p2_card.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "SENIOR CITIZENS & FAMILIES"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Targeted by aggressive digital arrest threats and fake police notices.\n• Pressured by claims of bank account freezes or legal seizures.\n• Needs high-contrast readability, spoken voice briefings, and 1-click 1930 Cyber Helpline access.\n• Needs clear step-by-step reassurance."
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    p3 = create_card(slide3, Inches(8.88), Inches(1.4), card_w, card_h, WHITE, CARD_BORDER)
    tf = p3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "DIGITAL BANKING BEGINNERS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Unfamiliar with official regulatory registries (SEBI SCORES, RBI Sachet).\n• Panicked by countdown timers (e.g. 'KYC expires in 15 mins').\n• Needs deterministic safety verification and calm containment protocols.\n• Needs explicit 'Safe Next Steps' rail."
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Bottom highlight bar
    hbar = create_card(slide3, Inches(0.8), Inches(5.9), Inches(11.73), Inches(0.8), LIGHT_BLUE, BORDER_BLUE)
    tf = hbar.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "BHARAT-FIRST INCLUSIVITY PILLAR"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "No login required • In-memory privacy • Bilingual English/Hindi toggle • Spoken audio briefings for non-readers."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    add_footer(slide3, 3)

    # ==========================================
    # SLIDE 4: Core Philosophy
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "CORE PHILOSOPHY", "Evidence First, Verdict Second", "Transforming ambiguous suspicion into objective regulatory proof before funds are transferred")

    # Left Card (35% width)
    lcard = create_card(slide4, Inches(0.8), Inches(1.4), Inches(4.2), Inches(5.2), LIGHT_BLUE, BORDER_BLUE)
    tf = lcard.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "THE PRE-ACTION BUFFER"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\nScamShield Bharat is strictly a Safety & Evidence Verification Console.\n\n• NOT an Investment Advisor\n• NO Stock Tips or Buy/Sell Calls\n• NO Black-Box Guesses\n\nInstead, it provides an institutional barrier that cross-examines claims directly against official regulatory circulars."
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Right 6-Step Pipeline (65% width)
    rcard = create_card(slide4, Inches(5.2), Inches(1.4), Inches(7.33), Inches(5.2), WHITE, CARD_BORDER)
    tf = rcard.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "THE 6-STAGE DETERMINISTIC SAFETY PIPELINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK

    steps = [
        ("STAGE 01: Ingestion & PII Masking", "In-memory scanning scrubs PAN, Aadhaar, phone numbers, and UPI IDs."),
        ("STAGE 02: Claim Extraction", "Isolates promised returns, urgency deadlines, and entity claims."),
        ("STAGE 03: Deterministic Risk Engine", "20 calibrated heuristic rules in English & Devanagari Hindi."),
        ("STAGE 04: Regulatory Cross-Exam", "Grounds claims in SEBI, RBI Sachet, MCA, and CERT-In advisories."),
        ("STAGE 05: Epistemic Uncertainty", "Explicitly discloses what could NOT be verified independently."),
        ("STAGE 06: Incident Response Protocol", "Calm 5-step containment checklist + 1-click 1930 Cyber Helpline.")
    ]
    for stitle, sdesc in steps:
        p_s = tf.add_paragraph()
        p_s.text = f"• {stitle}: {sdesc}"
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = SLATE

    add_footer(slide4, 4)

    # ==========================================
    # SLIDE 5: Real User Journey
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "USER JOURNEY", "One Message. One Safety Check.", "A seamless 3-second journey from receiving an unverified tip to taking safe defensive action")

    # 4 Horizontal Step Cards
    step_w = Inches(2.75)
    step_h = Inches(5.2)

    sdata = [
        ("1. INGESTION", DEEP_BLUE, LIGHT_BLUE, BORDER_BLUE, [
            "User receives suspicious tip on WhatsApp/Telegram/SMS.",
            "Pastes text, uploads screenshot, or submits web link.",
            "No account creation or password needed.",
            "Instant client-side payload validation."
        ]),
        ("2. PII MASKING", DEEP_BLUE, LIGHT_BLUE, BORDER_BLUE, [
            "Client & server regex scrubs sensitive details.",
            "PAN, Aadhaar, UPI IDs, phone numbers masked.",
            "0 disk writes, 0 database retention.",
            "Guaranteed privacy-first analysis."
        ]),
        ("3. CROSS-EXAM", RED_TEXT, RED_BG, RED_BORDER, [
            "Deterministic heuristics match 20 threat patterns.",
            "Correlates claims with official SEBI/RBI circulars.",
            "Generates 5-column signature Evidence Matrix.",
            "Calculates executive risk verdict."
        ]),
        ("4. ACTION RAIL", GREEN_TEXT, GREEN_BG, GREEN_BORDER, [
            "Bilingual executive safety briefing (EN/HI).",
            "Spoken voice readout via Web Speech API.",
            "5-step containment action protocol.",
            "1-click 1930 Cyber Helpline call."
        ])
    ]

    for idx, (stitle, tcolor, bgcolor, bdrcolor, bullets) in enumerate(sdata):
        left_pos = Inches(0.8 + idx * 3.0)
        sc = create_card(slide5, left_pos, Inches(1.4), step_w, step_h, bgcolor, bdrcolor)
        tf = sc.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = tcolor
        
        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"\n• {b}"
            pb.font.size = Pt(11)
            pb.font.color.rgb = SLATE

    add_footer(slide5, 5)

    # ==========================================
    # SLIDE 6: What ScamShield Detects
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "THREAT VECTORS", "Risk Signals That Matter Before Action", "20 deterministic rules calibrated against active financial fraud modus operandi in Bharat")

    card_w = Inches(2.78)
    card_h = Inches(2.5)

    threats = [
        ("GUARANTEED RETURNS", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Promises of assured 300% monthly profit or zero-risk stock returns prohibited by SEBI."),
        ("ARTIFICIAL URGENCY", "HIGH", AMBER_TEXT, AMBER_BG, AMBER_BORDER, "Coercive countdowns ('expires in 15 mins', '2 slots left') forcing hasty decisions."),
        ("REGULATORY IMPERSONATION", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Forged SEBI, RBI, MCA, or Police credentials, logos, and digital arrest threats."),
        ("MULE UPI ACCOUNTS", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Directing money into personal UPI handles or third-party savings accounts."),
        ("TELEGRAM/WHATSAPP FUNNELS", "HIGH", AMBER_TEXT, AMBER_BG, AMBER_BORDER, "Diverting users into unmonitored private VIP groups for illegal tip-sharing."),
        ("MALICIOUS APK SIDELOADING", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Urging installation of sideloaded apps or AnyDesk/QuickSupport tools."),
        ("FAKE KYC / ACCOUNT BLOCK", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Coercive claims that Demat or bank account is frozen to steal credentials."),
        ("CREDENTIAL / OTP DEMANDS", "CRITICAL", RED_TEXT, RED_BG, RED_BORDER, "Direct requests for trading PINs, netbanking passwords, or SMS OTPs.")
    ]

    for idx, (title, severity, tcolor, bgcolor, bdrcolor, desc) in enumerate(threats):
        row = idx // 4
        col = idx % 4
        left = Inches(0.8 + col * 2.98)
        top = Inches(1.4 + row * 2.65)
        
        c = create_card(slide6, left, top, card_w, card_h, bgcolor, bdrcolor)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = f"[{severity}]"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = tcolor
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(11.5)
        p2.font.bold = True
        p2.font.color.rgb = DARK

        p3 = tf.add_paragraph()
        p3.text = f"\n{desc}"
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = SLATE

    add_footer(slide6, 6)

    # ==========================================
    # SLIDE 7: Evidence-First Differentiation
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "DIFFERENTIATION", "Not “AI Says Scam.” Show The Evidence.", "Replacing black-box probabilistic guesswork with verifiable regulatory citations")

    col_w = Inches(5.7)
    col_h = Inches(5.2)

    # Left: Standard Black-Box AI
    c_left = create_card(slide7, Inches(0.8), Inches(1.4), col_w, col_h, CARD_BG, CARD_BORDER)
    tf = c_left.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "STANDARD BLACK-BOX AI CHATBOTS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED_TEXT

    left_points = [
        ("Hallucinated Confidence Scores", "Outputs arbitrary numbers (e.g. '98.4% Scam') without audit trail."),
        ("No Regulatory Citations", "Gives vague AI opinions without linking to official SEBI/RBI rules."),
        ("Binary Guesswork", "Forces binary true/false classifications even on incomplete context."),
        ("Security Risk", "Susceptible to prompt injection and malicious jailbreaks."),
        ("Passive Output", "Leaves user confused with no structured incident containment.")
    ]
    for title, desc in left_points:
        p = tf.add_paragraph()
        p.text = f"\n✗ {title}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    # Right: ScamShield Bharat
    c_right = create_card(slide7, Inches(6.83), Inches(1.4), col_w, col_h, LIGHT_BLUE, BORDER_BLUE)
    tf = c_right.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "SCAMSHIELD BHARAT EVIDENCE CONSOLE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    right_points = [
        ("Signature 5-Column Matrix", "Connects every claim to official regulatory source & explanation."),
        ("Grounding in Trusted Sources", "Direct clickable citations to sebi.gov.in, scores.gov.in, sachet.rbi.org.in."),
        ("Epistemic Honesty", "Explicitly surfaces 'What Could Not Be Verified' rather than guessing."),
        ("Hardened Security", "Deterministic risk engine immune to prompt injection jailbreaks."),
        ("Active Containment", "5-step defensive protocol with 1-click 1930 Cyber Helpline rail.")
    ]
    for title, desc in right_points:
        p = tf.add_paragraph()
        p.text = f"\n✓ {title}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    add_footer(slide7, 7)

    # ==========================================
    # SLIDE 8: Three Realistic Scenarios
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "TESTED SCENARIOS", "Three Real-World Attack Paths in Bharat", "Deterministic detection matching active hackathon evaluation cases in under 3 seconds")

    card_w = Inches(3.64)
    card_h = Inches(5.2)

    scenarios = [
        ("SCENARIO A: GUARANTEED RETURN", RED_TEXT, RED_BG, RED_BORDER, [
            ("Suspicious Message", "“Invest ₹10,000 in VIP Trading group, get ₹50,000 guaranteed weekly. Only 20 seats left!”"),
            ("Matched Threat Signals", "• Guaranteed Return Promise\n• Artificial Urgency Countdown\n• Unregistered Telegram Funnel"),
            ("Official Regulatory Evidence", "SEBI Advisory on Guaranteed Return Schemes (EVID-SEBI-GUARANTEED-RETURNS)"),
            ("Verdict", "HIGH CONCERN (Prohibited Scheme)")
        ]),
        ("SCENARIO B: FAKE KYC EXPIRY", AMBER_TEXT, AMBER_BG, AMBER_BORDER, [
            ("Suspicious Message", "“Urgent: Your Trading Demat Account will be blocked tonight. Update KYC: bit.ly/kyc-verify”"),
            ("Matched Threat Signals", "• Account Suspension Threat\n• Shortened Phishing URL\n• Fake KYC Portal Clone"),
            ("Official Regulatory Evidence", "RBI Sachet Advisory on Phishing SMS & Clone Portals (EVID-RBI-SACHET-UNREGISTERED)"),
            ("Verdict", "HIGH CONCERN (Phishing Attack)")
        ]),
        ("SCENARIO C: REMOTE APK SIDELOAD", RED_TEXT, RED_BG, RED_BORDER, [
            ("Suspicious Message", "“Bank support officer: Download Fast-KYC-Support.apk to unblock failed UPI transaction.”"),
            ("Matched Threat Signals", "• Banking Authority Impersonation\n• Sideloaded Android APK\n• Remote Access Trojan Threat"),
            ("Official Regulatory Evidence", "CERT-In Advisory on Malicious Banking APKs & Screen Sharing (EVID-CERTIN-MALICIOUS-APKS)"),
            ("Verdict", "HIGH CONCERN (Device Takeover)")
        ])
    ]

    for idx, (stitle, tcolor, bgcolor, bdrcolor, details) in enumerate(scenarios):
        left = Inches(0.8 + idx * 4.04)
        c = create_card(slide8, left, Inches(1.4), card_w, card_h, bgcolor, bdrcolor)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = tcolor

        for heading, body in details:
            ph = tf.add_paragraph()
            ph.text = f"\n{heading}:"
            ph.font.size = Pt(10.5)
            ph.font.bold = True
            ph.font.color.rgb = DARK
            
            pb = tf.add_paragraph()
            pb.text = body
            pb.font.size = Pt(10)
            pb.font.color.rgb = SLATE

    add_footer(slide8, 8)

    # ==========================================
    # SLIDE 9: Technology Stack & Decision Engine
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "ARCHITECTURE", "Hybrid Architecture: Deterministic Where It Matters", "Robust engineering combining deterministic heuristics with generative AI explanations")

    col_w = Inches(3.64)
    col_h = Inches(4.3)

    # Pillar 1
    p1 = create_card(slide9, Inches(0.8), Inches(1.4), col_w, col_h, WHITE, CARD_BORDER)
    tf = p1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "FRONTEND & EDGE RUNTIME"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Next.js 14 (App Router)\n• React 18 with Strict Mode\n• Tailwind CSS Design System\n• Vercel Serverless Edge Platform\n• Responsive Desktop & Mobile\n• Native Web Speech API"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Pillar 2
    p2_card = create_card(slide9, Inches(4.84), Inches(1.4), col_w, col_h, WHITE, CARD_BORDER)
    tf = p2_card.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "DETERMINISTIC RISK ENGINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• TypeScript Strict Type Safety\n• Zod v3 Runtime Validation\n• In-Memory Regex Pattern Matcher\n• Bilingual English & Hindi Stems\n• In-Memory Evidence Registry\n• 0 Disk Storage Requirement"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Pillar 3
    p3 = create_card(slide9, Inches(8.88), Inches(1.4), col_w, col_h, WHITE, CARD_BORDER)
    tf = p3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "AI & VOICE SYNTHESIS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Google Gemini 1.5 Flash\n• Plain-Language Summary Generator\n• Zero Hallucination Guardrails\n• Client & Server PII Redaction\n• Bilingual Voice Synthesis\n• Sub-3s Analysis Response Time"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE

    # Bottom Stats Bar
    sbar = create_card(slide9, Inches(0.8), Inches(5.9), Inches(11.73), Inches(0.8), LIGHT_BLUE, BORDER_BLUE)
    tf = sbar.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "ENGINEERING BENCHMARKS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "79/79 Vitest Tests Passing • 54/54 Red-Team Attack Vectors Cleared • 0 Security Vulnerabilities in Git Audit."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    add_footer(slide9, 9)

    # ==========================================
    # SLIDE 10: Bharat-First Design
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "INCLUSIVITY", "Designed for Bharat, Not Just Expert Users", "Breaking language, literacy, and technical barriers for inclusive financial safety")

    card_w = Inches(5.66)
    card_h = Inches(2.55)

    # Card 1: Bilingual
    b1 = create_card(slide10, Inches(0.8), Inches(1.4), card_w, card_h, WHITE, CARD_BORDER)
    tf = b1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "1. Full Bilingual English & Hindi (स्कैमशील्ड भारत)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Instant 1-click language switch across all screens and evidence stages.\n• Natural Devanagari Hindi translations for all technical findings.\n• Eliminates alienating English-only financial jargon."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Card 2: Voice Briefing
    b2 = create_card(slide10, Inches(6.86), Inches(1.4), card_w, card_h, WHITE, CARD_BORDER)
    tf = b2.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "2. Audio Briefing Voice Reader"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• Integrated Web Speech API reads risk summaries and safe actions aloud.\n• Empowers non-literate and visually impaired investors.\n• Clear spoken Hindi and English narration."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Card 3: Low Cognitive Load
    b3 = create_card(slide10, Inches(0.8), Inches(4.15), card_w, card_h, WHITE, CARD_BORDER)
    tf = b3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "3. Low Cognitive Load Design"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n• High-contrast institutional color coding (High Concern, Caution, Verified).\n• Clear 4-metric statistics and scannable Stage Headings.\n• Replaces anxiety with structured, actionable clarity."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    # Card 4: Action Rail
    b4 = create_card(slide10, Inches(6.86), Inches(4.15), card_w, card_h, WHITE, CARD_BORDER)
    tf = b4.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.22)
    p = tf.paragraphs[0]
    p.text = "4. Direct 1930 Cyber Helpline Action Rail"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_TEXT
    p2 = tf.add_paragraph()
    p2.text = "\n• 1-click direct dial to National Cyber Crime Helpline 1930.\n• Clickable access to official cybercrime.gov.in & SEBI SCORES portals.\n• 5-step containment checklist: Stop, Protect, Verify, Report, Recover."
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    add_footer(slide10, 10)

    # ==========================================
    # SLIDE 11: Security & Red-Team Verification
    # ==========================================
    slide11 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide11, "SECURITY AUDIT", "Security Is Built Into The Product", "Comprehensive verification against prompt injection, SSRF, XSS, and data leakage")

    col_w = Inches(5.7)
    col_h = Inches(5.2)

    # Left: Security Controls
    c_left = create_card(slide11, Inches(0.8), Inches(1.4), col_w, col_h, WHITE, CARD_BORDER)
    tf = c_left.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "ARCHITECTURAL SECURITY CONTROLS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    sec_controls = [
        ("Zero Data Retention", "In-memory processing only; raw user messages discarded immediately."),
        ("Automated PII Scrubbing", "Regex filters mask Aadhaar, PAN, phone numbers, and UPI IDs."),
        ("Anti-SSRF URL Engine", "Blocks loopback (127.0.0.1), private RFC 1918, octal/hex, & AWS/GCP metadata IP blocks."),
        ("Upload Hardening", "5MB file ceiling with magic-byte file signature validation (JPG/PNG/WebP)."),
        ("DoS & Rate Limiting", "50KB payload limit + sliding-window rate limit (30 req/min/IP).")
    ]
    for title, desc in sec_controls:
        p = tf.add_paragraph()
        p.text = f"\n• {title}:\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    # Right: Red Team Scorecard
    c_right = create_card(slide11, Inches(6.83), Inches(1.4), col_w, col_h, LIGHT_BLUE, BORDER_BLUE)
    tf = c_right.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "RED-TEAM SECURITY AUDIT SCORECARD"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_TEXT

    audit_items = [
        ("Automated Test Suite", "79 / 79 PASS (`vitest run`) across 7 test suites."),
        ("Localhost Red-Team Suite", "54 / 54 PASS (`http://localhost:3000`) across 13 threat categories."),
        ("Production Red-Team Suite", "54 / 54 PASS (`scamshield-bharat-kappa.vercel.app`)."),
        ("Prompt Injection Defense", "100% Protected (English & Devanagari system prompt extraction attempts neutralized)."),
        ("Secret Leakage Audit", "0 Secrets Found (Clean Git tree & environment variable isolation).")
    ]
    for title, desc in audit_items:
        p = tf.add_paragraph()
        p.text = f"\n✓ {title}:\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    add_footer(slide11, 11)

    # ==========================================
    # SLIDE 12: Trust & Guardrails
    # ==========================================
    slide12 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide12, "COMPLIANCE", "A Safety Tool — Not An Investment Advisor", "Strict boundary compliance adhering to SEBI (Investment Advisers) Regulations, 2013")

    col_w = Inches(5.7)
    col_h = Inches(5.2)

    # Left: Prohibited Actions
    c_left = create_card(slide12, Inches(0.8), Inches(1.4), col_w, col_h, RED_BG, RED_BORDER)
    tf = c_left.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "STRICTLY PROHIBITED (NON-ADVISORY BOUNDARY)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RED_TEXT

    prohibited = [
        ("NO Stock Tips or Recommendations", "ScamShield never tells users what stocks or securities to buy or sell."),
        ("NO Buy / Sell / Hold Calls", "No investment advisory or trade signals are generated."),
        ("NO Price Targets or Projections", "Never predicts market movements, returns, or asset values."),
        ("NO Automated Trading Management", "Does not execute trades or manage user funds."),
        ("NO Broker Affiliation", "Does not promote or favor any specific commercial broker.")
    ]
    for title, desc in prohibited:
        p = tf.add_paragraph()
        p.text = f"\n✗ {title}\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    # Right: Permitted & Delivered
    c_right = create_card(slide12, Inches(6.83), Inches(1.4), col_w, col_h, GREEN_BG, GREEN_BORDER)
    tf = c_right.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "PERMITTED & DELIVERED (SAFETY BUFFER)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_TEXT

    delivered = [
        ("YES Objective Risk Signal Extraction", "Flags guaranteed returns, urgency deadlines, and suspicious funnels."),
        ("YES Regulatory Cross-Examination", "Direct citations to SEBI, RBI, MCA, and CERT-In advisories."),
        ("YES Epistemic Uncertainty Disclosures", "Explicitly states what claims could not be verified."),
        ("YES Incident Response Guidance", "Clear 5-step containment checklist to protect compromised accounts."),
        ("YES Official Helplines", "Direct links to National Cybercrime Helpline 1930 & SEBI SCORES.")
    ]
    for title, desc in delivered:
        p = tf.add_paragraph()
        p.text = f"\n✓ {title}\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    add_footer(slide12, 12)

    # ==========================================
    # SLIDE 13: Impact & Scalability
    # ==========================================
    slide13 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide13, "SCALABILITY", "From One Suspicious Message to a National Safety Layer", "Clear separation between verified live prototype and future strategic roadmap")

    col_w = Inches(5.7)
    col_h = Inches(5.2)

    # Left: Built & Operational Now
    c_left = create_card(slide13, Inches(0.8), Inches(1.4), col_w, col_h, LIGHT_BLUE, BORDER_BLUE)
    tf = c_left.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "LIVE PROTOTYPE (BUILT & OPERATIONAL NOW)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    live_features = [
        ("Global Edge Deployment", "Live on Vercel Edge CDN with sub-3s response latency."),
        ("Multimodal Input Console", "Supports text messages, image screenshots, and web URLs."),
        ("Bilingual English & Hindi", "Full UI localization and voice briefing support."),
        ("Deterministic Heuristic Engine", "20 risk rules and in-memory regulatory cross-exam matrix."),
        ("Hardened Security Suite", "79 automated tests and 54 red-team vectors validated.")
    ]
    for title, desc in live_features:
        p = tf.add_paragraph()
        p.text = f"\n• {title}:\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    # Right: Future Expansion
    c_right = create_card(slide13, Inches(6.83), Inches(1.4), col_w, col_h, WHITE, CARD_BORDER)
    tf = c_right.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "STRATEGIC FUTURE ROADMAP (PLANNED)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK

    roadmap = [
        ("Regional Language Expansion", "Adding voice & text support for Tamil, Telugu, Bengali, Marathi, and Gujarati."),
        ("WhatsApp & Telegram Tipline Bot", "Direct message forwarding to verified ScamShield verification bots."),
        ("Direct SEBI Registry API Sync", "Live real-time querying of SEBI registered intermediary databases."),
        ("Distributed Redis Rate Limiting", "Upstash Redis backing for horizontal multi-million request scale."),
        ("Offline PWA Support", "Client-side caching of top 100 scam patterns for offline safety.")
    ]
    for title, desc in roadmap:
        p = tf.add_paragraph()
        p.text = f"\n• {title}:\n  {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE

    add_footer(slide13, 13)

    # ==========================================
    # SLIDE 14: Live Demo Walkthrough
    # ==========================================
    slide14 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide14, "LIVE DEMO", "See It Work — Live Production Deployment", "Verify end-to-end performance on desktop or mobile in real time")

    # Center Hero Card
    demo_card = create_card(slide14, Inches(0.8), Inches(1.4), Inches(11.73), Inches(5.2), WHITE, CARD_BORDER)
    tf = demo_card.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "RECOMMENDED JUDGING DEMO SEQUENCE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    steps = [
        ("Step 1: Open Live Application", "Visit https://scamshield-bharat-kappa.vercel.app/ on desktop or mobile."),
        ("Step 2: Load Deterministic Demo", "Click 'Live Demo' or choose 'Guaranteed Return' / 'Fake KYC' / 'APK Sideload'."),
        ("Step 3: Inspect Evidence Trail", "Observe the 5-column cross-examination matrix connecting claims to SEBI advisories."),
        ("Step 4: Toggle Bilingual Mode", "Click 'हिंदी' to verify complete Devanagari translation of analysis and evidence."),
        ("Step 5: Test Audio Briefing", "Click 'Listen Audio' to hear native voice readout of safety findings."),
        ("Step 6: Test Custom Message / Screenshot", "Paste your own suspicious message or upload a forward screenshot."),
        ("Step 7: Verify Safe Action Rail", "Review the 5 containment steps and direct click-to-call 1930 Cyber Helpline.")
    ]
    for stitle, sdesc in steps:
        p = tf.add_paragraph()
        p.text = f"• {stitle}: {sdesc}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = SLATE

    add_footer(slide14, 14)

    # ==========================================
    # SLIDE 15: Why ScamShield Fits SANGYAN
    # ==========================================
    slide15 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide15, "SANGYAN ALIGNMENT", "Built Around Investor Resilience", "Empowering Indian citizens with an evidence-first digital defense shield")

    card_w = Inches(5.66)
    card_h = Inches(2.2)

    pillars = [
        ("1. PRE-TRANSACTION RESILIENCE", DEEP_BLUE, LIGHT_BLUE, BORDER_BLUE, "Prevents catastrophic capital loss at the critical decision point before money leaves."),
        ("2. BHARAT-FIRST ACCESSIBILITY", DEEP_BLUE, LIGHT_BLUE, BORDER_BLUE, "Bilingual Hindi/English, spoken voice briefing, low cognitive load, and mobile-first responsiveness."),
        ("3. STRICT TRUST & GUARDRAILS", GREEN_TEXT, GREEN_BG, GREEN_BORDER, "Strict non-advisory boundary; in-memory PII protection; explicit uncertainty disclosure."),
        ("4. ENGINEERING RIGOR", GREEN_TEXT, GREEN_BG, GREEN_BORDER, "79/79 automated tests, 54/54 red-team vectors cleared, 0 secret leaks, live Vercel production deployment.")
    ]

    for idx, (title, tcolor, bgcolor, bdrcolor, desc) in enumerate(pillars):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.06)
        top = Inches(1.4 + row * 2.35)

        c = create_card(slide15, left, top, card_w, card_h, bgcolor, bdrcolor)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = tcolor
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE

    # Grand Closing Callout Banner
    gban = create_card(slide15, Inches(0.8), Inches(6.15), Inches(11.73), Inches(0.65), DEEP_BLUE, DEEP_BLUE)
    tf = gban.text_frame
    tf.margin_left = tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "SCAMSHIELD BHARAT — “Check the claim before you act.”  |  Live: scamshield-bharat-kappa.vercel.app"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    add_footer(slide15, 15)

    # Save presentation
    prs.save(output_path)
    print(f"Successfully generated clean white PowerPoint presentation at: {output_path}")

if __name__ == '__main__':
    out_file = os.path.join(os.getcwd(), 'presentation', 'SCAMSHIELD-BHARAT-FINAL-SANGYAN-DECK.pptx')
    build_presentation(out_file)
